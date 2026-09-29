# -*- coding: utf-8 -*-
"""动态内容评分规则 → Excel（两个 sheet：内容基础分 / 时间衰减系数）"""
try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def xl_color(css_hex):
    value = css_hex[1:] if css_hex.startswith("#") else css_hex
    value = value.upper()
    if len(value) != 6:
        raise ValueError("Expected #RRGGBB, got: %s" % css_hex)
    return "FF" + value


XL_HEADER_BG = xl_color("#4472C4")
XL_HEADER_FG = xl_color("#FFFFFF")
XL_TOTAL_BG = xl_color("#2F5597")
XL_SUBTITLE_FG = xl_color("#2F5597")
XL_LIGHT = xl_color("#D9E2F3")
XL_BORDER = xl_color("#BFBFBF")
thin = Side(style="thin", color=XL_BORDER)
BORDER_ALL = Border(left=thin, right=thin, top=thin, bottom=thin)

HEADER_FONT = Font(bold=True, color=XL_HEADER_FG)
HEADER_FILL = PatternFill("solid", fgColor=XL_HEADER_BG)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
TOTAL_FILL = PatternFill("solid", fgColor=XL_TOTAL_BG)
TOTAL_FONT = Font(bold=True, color="FFFFFFFF")
SUBTITLE_FONT = Font(bold=True, size=12, color=XL_SUBTITLE_FG)

wb = openpyxl.Workbook()
wb.properties.title = "动态内容评分规则"

# ───────────────────────── Sheet1 内容基础分 ─────────────────────────
ws = wb.active
ws.title = "内容基础分"

# 锚点：A1 大标题 / A2 副标题 / A3:C3 表头 / A4:C6 维度数据 / A7 合计 / A9 小标题 / A10:B10 表头 / A11:B16 文字档位
ws["A1"] = "内容基础分（上限合计 97 分）"
ws["A1"].font = Font(bold=True, size=14)
ws.merge_cells("A1:C1")

ws["A2"] = "一、各维度权重"
ws["A2"].font = SUBTITLE_FONT
ws.merge_cells("A2:C2")

for col, name in zip("ABC", ["内容维度", "计分规则", "上限分数"]):
    ws[f"{col}3"] = name
    ws[f"{col}3"].font = HEADER_FONT
    ws[f"{col}3"].fill = HEADER_FILL
    ws[f"{col}3"].alignment = HEADER_ALIGN
    ws[f"{col}3"].border = BORDER_ALL

# 数据行：文字（6 档合并 4~9 行？——改为规则摘要 + 明细表方案：数据 4~6 行，合计 7 行）
ws["A4"] = "文字"
ws["B4"] = "按字符数分 6 档计分，明细见下方「文字档位明细」"
ws["C4"] = 50
ws["A5"] = "图片"
ws["B5"] = "一张图片 = 5 分"
ws["C5"] = 45
ws["A6"] = "歌曲"
ws["B6"] = "一首歌曲 = 2 分"
ws["C6"] = 2
for r in range(4, 7):
    for col in "ABC":
        c = ws[f"{col}{r}"]
        c.border = BORDER_ALL
        c.alignment = CENTER if col != "B" else LEFT_WRAP
    ws[f"C{r}"].number_format = "0"
    ws[f"C{r}"].alignment = Alignment(horizontal="center", vertical="center")

ws["A7"] = "合计"
ws["B7"] = "累计加分"
ws["C7"] = "=SUM(C4:C6)"
for col in "ABC":
    c = ws[f"{col}7"]
    c.font = TOTAL_FONT
    c.fill = TOTAL_FILL
    c.border = BORDER_ALL
    c.alignment = Alignment(horizontal="center", vertical="center")
ws["C7"].number_format = "0"

# 文字档位明细（锚点：A9 小标题 / A10:B10 表头 / A11:B16 数据）
ws["A9"] = "文字档位明细"
ws["A9"].font = SUBTITLE_FONT
ws.merge_cells("A9:B9")
for col, name in zip("AB", ["字符数区间", "得分"]):
    ws[f"{col}10"] = name
    ws[f"{col}10"].font = HEADER_FONT
    ws[f"{col}10"].fill = HEADER_FILL
    ws[f"{col}10"].alignment = HEADER_ALIGN
    ws[f"{col}10"].border = BORDER_ALL
TIERS = [("0-10 个字符", 5), ("11-50 个字符", 10), ("51-100 个字符", 20),
         ("101-150 个字符", 30), ("151-200 个字符", 40), ("200 个字符以上", 50)]
for i, (rng, score) in enumerate(TIERS):
    r = 11 + i
    ws[f"A{r}"] = rng
    ws[f"B{r}"] = score
    ws[f"A{r}"].border = BORDER_ALL
    ws[f"B{r}"].border = BORDER_ALL
    ws[f"A{r}"].alignment = CENTER
    ws[f"B{r}"].alignment = CENTER
    ws[f"B{r}"].number_format = "0"

ws.column_dimensions["A"].width = 16
ws.column_dimensions["B"].width = 46
ws.column_dimensions["C"].width = 12

# ───────────────────────── Sheet2 时间衰减系数 ─────────────────────────
ws2 = wb.create_sheet("时间衰减系数")

ws2["A1"] = "时间衰减系数（基于动态发布时间）"
ws2["A1"].font = Font(bold=True, size=14)
ws2.merge_cells("A1:B1")

ws2["A2"] = "加入时间衰减系数"
ws2["A2"].font = SUBTITLE_FONT
ws2.merge_cells("A2:B2")

# 锚点：A3:B3 表头 / A4:B8 数据 / A9 备注
for col, name in zip("AB", ["动态发布时间", "时间系数"]):
    ws2[f"{col}3"] = name
    ws2[f"{col}3"].font = HEADER_FONT
    ws2[f"{col}3"].fill = HEADER_FILL
    ws2[f"{col}3"].alignment = HEADER_ALIGN
    ws2[f"{col}3"].border = BORDER_ALL

DECAY = [("0～24 小时", 1.0), (">24～48 小时", 0.8), (">48～72 小时", 0.6),
         (">3～7 天", 0.4), (">7 天以上", 0.2)]
for i, (rng, factor) in enumerate(DECAY):
    r = 4 + i
    ws2[f"A{r}"] = rng
    ws2[f"B{r}"] = factor
    ws2[f"A{r}"].border = BORDER_ALL
    ws2[f"B{r}"].border = BORDER_ALL
    ws2[f"A{r}"].alignment = CENTER
    ws2[f"B{r}"].alignment = CENTER
    ws2[f"B{r}"].number_format = "0.0"

ws2["A9"] = "动态最终得分 = 内容基础分 × 时间系数"
ws2["A9"].font = Font(bold=True, color=XL_SUBTITLE_FG)
ws2.merge_cells("A9:B9")
ws2["A9"].alignment = LEFT_WRAP

ws2.column_dimensions["A"].width = 22
ws2.column_dimensions["B"].width = 14

OUT = "/Users/ouyang/Desktop/AutomationTest/动态内容评分规则.xlsx"
wb.save(OUT)
print("saved:", OUT)
